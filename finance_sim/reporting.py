"""Curated PNG and Excel snapshots derived from committed simulation results."""

from __future__ import annotations

import json

import matplotlib
import numpy as np
import pandas as pd

from pathlib import Path

from .results import paired_amanda_comparison
from .decisions import generate_decision_report
from .currency import real_ledger
from .storage import read_table, write_json
from display_names import display_name
from matplotlib.ticker import FuncFormatter, MaxNLocator


matplotlib.use("Agg")
import matplotlib.pyplot as plt


INK, GREEN, GOLD, PALE = "#233d38", "#397c69", "#b38335", "#f4f2e9"


def flatten(value: object, prefix: str = "") -> list[dict]:
    if isinstance(value, dict):
        return [row for key, child in value.items() for row in flatten(child, f"{prefix}.{key}".strip("."))]
    if isinstance(value, list):
        return [row for i, child in enumerate(value) for row in flatten(child, f"{prefix}[{i}]" )]
    return [{"setting": prefix, "value": value}]


def save_figure(fig: plt.Figure, folder: Path, number: int, name: str, description: str) -> dict:
    def dollars(value: float, position: int) -> str:
        amount = abs(value)
        suffix, divisor = ("B", 1e9) if amount >= 1e9 else (("M", 1e6) if amount >= 1e6 else (("k", 1e3) if amount >= 1e3 else ("", 1)))
        return ("-" if value < 0 else "") + f"${amount/divisor:g}{suffix}"
    for ax in fig.axes:
        if "USD" in ax.get_ylabel():
            ax.yaxis.set_major_formatter(FuncFormatter(dollars))
            ax.yaxis.set_major_locator(MaxNLocator(7))
        if "USD" in ax.get_xlabel() or "net worth difference" in ax.get_xlabel().lower():
            ax.xaxis.set_major_formatter(FuncFormatter(dollars))
            ax.xaxis.set_major_locator(MaxNLocator(7))
        ax.margins(x=.03, y=.08)
    fig.text(.01, .015, "Conditional model outcomes | 2026 USD unless labeled | Provisional assumptions apply", fontsize=8, color=INK)
    fig.tight_layout(rect=(0, .04, 1, 1))
    filename = f"{number:02d}_{name}.png"
    fig.savefig(folder / "figures" / filename, dpi=300, facecolor=PALE)
    plt.close(fig)
    return {"file": filename, "description": description, "dpi": 300}


def generate_reports(folder: Path, scenario_ids: list[str] | None = None) -> list[dict]:
    summary = read_table(folder, "summary", scenario_ids)
    if summary.empty:
        raise ValueError("No committed scenarios to report")
    # Keep all catalog summaries while bounding distribution-plot memory.
    ids = summary.iloc[np.linspace(0, len(summary)-1, min(1000, len(summary)), dtype=int)].scenario_id.tolist()
    terminal = read_table(folder, "terminal", ids)
    for column in list(terminal.columns):
        if column.startswith("real_") and column[5:] in terminal:
            terminal[column[5:]] = terminal[column]
    family_rows = summary[(summary.children > 0) & (summary.stop_after_birth == 0) & (summary.housing == "buy")]
    if family_rows.empty:
        family_rows = summary[(summary.children > 0) & (summary.stop_after_birth == 0)]
    first_id = str(family_rows.iloc[0].scenario_id) if not family_rows.empty else ids[0]
    monthly = read_table(folder, "monthly", [first_id])
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    report_scope = {"scenario_ids": ids, "reported_scenarios": len(summary), "distribution_scenarios": len(ids), "completed_scenarios": manifest["completed_scenarios"],
                    "selection": "All scenario summaries and risk combinations; Excel tables capped at 10,000 rows; distribution figures use at most 1,000 evenly spaced catalog scenarios",
                    "trajectory_scenario": first_id, "notice": "Mixture plots equally weight the selected conditional scenarios; they are not a probability forecast."}
    write_json(folder / "report_scope.json", report_scope)
    if manifest.get("dollar_basis") != "2026 USD":
        raise ValueError("Use the archived report engine for legacy nominal runs; do not relabel old results as real dollars")
    risks, frontier = generate_decision_report(folder, summary)
    summary.to_csv(folder / "scenario_summary.csv", index=False)
    paired = paired_amanda_comparison(summary)
    paired.to_csv(folder / "amanda_comparison.csv", index=False)
    plt.rcParams.update({"font.family": "DejaVu Sans", "axes.spines.top": False, "axes.spines.right": False,
                         "axes.facecolor": PALE, "figure.facecolor": PALE, "text.color": INK,
                         "axes.labelcolor": INK, "xtick.color": INK, "ytick.color": INK,
                         "axes.titleweight": "bold", "axes.grid": True, "grid.alpha": .16})
    chart_index = []
    trajectories = [("net_worth", "Net worth"), ("financial_net_worth", "Financial net worth"),
                    ("liquid_wealth", "Liquid wealth"), ("cash", "Cash reserves"),
                    ("retirement", "Retirement assets"), ("debt", "Loan balances"),
                    ("brokerage", "Taxable investments"), ("home_equity", "Home equity"),
                    ("earned_income", "Monthly earned income"), ("childcare", "Monthly paid childcare"),
                    ("health_premium", "Monthly health premiums"), ("tax_expense", "Monthly tax expense")]
    for number, (metric, title) in enumerate(trajectories, 1):
        frame = monthly[monthly.metric == metric].sort_values("date")
        fig, ax = plt.subplots(figsize=(9, 4.6))
        ax.fill_between(frame.date, frame.q025, frame.q975, color=GREEN, alpha=.12, label="95% predictive range")
        ax.fill_between(frame.date, frame.q100, frame.q900, color=GREEN, alpha=.20, label="80%")
        ax.fill_between(frame.date, frame.q250, frame.q750, color=GREEN, alpha=.3, label="50%")
        ax.plot(frame.date, frame.q500, color=GREEN, label="Median")
        ax.set(title=title, ylabel="USD", xlabel="Year")
        ax.ticklabel_format(axis="y", style="plain")
        ax.legend(fontsize=8, ncol=4)
        chart_index.append(save_figure(fig, folder, number, metric, f"Predictive bands for scenario {first_id}; marginal percentiles are not an individual ledger."))
    fig, ax = plt.subplots(figsize=(9, 4.6))
    ax.hist(terminal.net_worth, bins=40, color=GREEN)
    ax.set(title="Terminal net worth across selected scenarios", xlabel="USD", ylabel="Simulated outcomes")
    chart_index.append(save_figure(fig, folder, 13, "terminal_distribution", report_scope["notice"]))
    fig, ax = plt.subplots(figsize=(9, 4.6))
    ordered = np.sort(terminal.real_net_worth.to_numpy())  # (n,)
    ax.plot(ordered, np.arange(1, len(ordered)+1) / len(ordered), color=GREEN)
    ax.set(title="Purchasing power at the end of 2036", xlabel="2026 USD", ylabel="Cumulative share of outcomes")
    chart_index.append(save_figure(fig, folder, 14, "real_wealth_cdf", report_scope["notice"]))
    for number, metric, title in ((15, "shortfall_probability", "Probability of an unmet household bill"),
                                  (16, "buy_probability", "Probability of buying within the horizon")):
        fig, ax = plt.subplots(figsize=(9, 4.6))
        selected = summary.sort_values(metric, ascending=False).head(20)
        ax.barh(np.arange(len(selected)), selected[metric], color=GREEN)
        ax.set_yticks(np.arange(len(selected)), [f"{display_name(r.location)}, {r.children} kids, stop {r.stop_after_birth}" for r in selected.itertuples()], fontsize=7)
        ax.set(title=title, xlabel="Probability", xlim=(0, 1))
        chart_index.append(save_figure(fig, folder, number, metric, "Twenty highest modeled probabilities; exact binomial Monte Carlo intervals are in the summary table."))
    merged = terminal.merge(summary[["scenario_id", "career", "location"]], on="scenario_id")
    for number, category in ((17, "career"), (18, "location")):
        fig, ax = plt.subplots(figsize=(9, 4.6))
        groups = list(merged.groupby(category, sort=False))
        ax.boxplot([frame.net_worth.to_numpy() for _, frame in groups], tick_labels=[display_name(name) for name, _ in groups], showfliers=False)
        ax.tick_params(axis="x", labelrotation=25)
        ax.set(title=f"Terminal outcomes by {category}", ylabel="Net worth (USD)")
        chart_index.append(save_figure(fig, folder, number, category, "Grouping mixes other scenario assumptions; this is not a causal estimate."))
    fig, ax = plt.subplots(figsize=(9, 4.6))
    if not paired.empty:
        selected = paired.head(20)
        ax.barh(np.arange(len(selected)), selected.net_worth_difference, color=GOLD)
        ax.set_yticks(np.arange(len(selected)), [f"{display_name(r.location)}, stop after child {r.stop_after_birth_stay_home}" for r in selected.itertuples()], fontsize=7)
        ax.set(xlabel="Median net worth difference (stay home minus working)")
    else:
        ax.text(.5, .5, "No matched working / stay-at-home pair in this report selection", ha="center", transform=ax.transAxes)
    ax.set_title("Amanda working versus staying home")
    chart_index.append(save_figure(fig, folder, 19, "amanda_work_comparison", "Matched on all other scenario dimensions; negative values mean lower wealth when staying home."))
    # Select a financing/exit case for company figures instead of plotting a
    # no-exit baseline with empty dilution and waterfall charts.
    company_rows = summary[(summary.exit_value > 0) & (summary.company_outcome == "vc")]
    company_id = str(company_rows.iloc[0].scenario_id) if not company_rows.empty else first_id
    report_scope["company_ledger_scenario"] = company_id
    ledgers = sorted((folder / "ledgers").glob(f"{company_id}-*.parquet"))
    if ledgers:
        ledger = pd.read_parquet(ledgers[len(ledgers)//2])
    else:
        from .configuration import load_config
        from .engine import simulate
        from .results import SimulationResult
        from .scenarios import scenario_from_record
        config = load_config(folder / "config.json")
        scenario = scenario_from_record(summary[summary.scenario_id == company_id].iloc[0].to_dict())
        simulation = SimulationResult(*simulate(config, scenario))
        ordered = np.argsort(simulation.series("net_worth")[:, -1])
        path = int(ordered[len(ordered)//2])
        ledger = simulation.ledger(path, config)
        ledger.to_parquet(folder / "ledgers" / f"{company_id}-{path}.parquet", index=False)
    write_json(folder / "report_scope.json", report_scope)
    ledger = real_ledger(ledger)
    for number, columns, title in ((20, ["business_revenue", "business_grant_receipts", "business_financing"], "Company revenue and funding"),
                                   (21, ["business_company_cash", "business_company_arrears", "business_restricted_cash"], "Company liquidity and unpaid compensation"),
                                   (22, ["business_founder_share"], "Founder ownership through financing"),
                                   (23, ["business_exit_gross", "exit_cash", "business_advisor_payout"], "Exit proceeds and payment timing"),
                                   (24, ["earned_income", "living_cost", "childcare", "health_premium", "tax_expense"], "Cumulative household income and costs")):
        fig, ax = plt.subplots(figsize=(9, 4.6))
        if not ledger.empty:
            for column in columns:
                values = ledger[column].cumsum() if number == 24 else ledger[column]
                ax.plot(ledger.date, values, label=column.replace("business_", "").replace("_", " "))
            ax.legend(fontsize=8)
        else:
            ax.text(.5, .5, "Regenerate a selected path ledger to inspect this schedule", ha="center", transform=ax.transAxes)
        ax.set(title=title, xlabel="Year", ylabel="Ownership fraction" if number == 22 else "USD")
        chart_index.append(save_figure(fig, folder, number, title.lower().replace(" ", "_"), "One retained, actual simulated path. This path is not the median of each component."))
    if "minimum_retirement_age" in terminal:
        fig, ax = plt.subplots(figsize=(9, 4.6))
        ages = terminal.minimum_retirement_age.dropna()
        if len(ages):
            ax.hist(ages, bins=np.arange(20, 86, 2), color=GREEN)
        ax.set(title="Earliest retirement age, including projections", xlabel="Logan's age", ylabel="Simulated outcomes")
        ax.text(.98, .96, f"Not reached by configured maximum: {terminal.minimum_retirement_age.isna().mean():.1%}",
                ha="right", va="top", transform=ax.transAxes, fontsize=9)
        chart_index.append(save_figure(fig, folder, 25, "minimum_retirement_age", "Selected conditional paths; beyond-horizon ages are projections. Unreached paths are reported separately, never dropped from scenario retirement-age quantiles."))
    if not risks.empty:
        fig, ax = plt.subplots(figsize=(9, 4.6))
        highest = risks.head(12)
        ax.barh(np.arange(len(highest)), highest.distress_probability, color=GOLD)
        ax.set_yticks(np.arange(len(highest)), [f"{display_name(r.location)}, {r.births or 'no children'}, {display_name(r.macro)}" for r in highest.itertuples()], fontsize=7)
        ax.set(title="Combinations with the highest liquidity risk", xlabel="Probability of depleted funds or an unmet bill", xlim=(0, 1))
        chart_index.append(save_figure(fig, folder, 26, "depletion_combinations", "Highest-risk completed conditional scenarios; full choices and Monte Carlo bounds are in depletion_combinations.csv."))
    if not frontier.empty:
        fig, ax = plt.subplots(figsize=(9, 4.6))
        for family, frame in frontier[(frontier.macro == "all_tested_macros") & (frontier.location == "all_tested_locations")].groupby("salary_family"):
            midpoint = (frame.salary_low_2026 + frame.salary_high_2026) / 2
            ax.plot(midpoint, frame.risk_estimate, marker="o", label=display_name(family))
            ax.plot(midpoint, frame.risk_upper_simultaneous_95, linestyle=":", alpha=.7, label=f"{display_name(family)} upper bound")
        ax.set(title="Salary and risk across all tested life choices", xlabel="Starting salary (2026 USD)", ylabel="Any-choice distress probability", ylim=(-.03, 1.03))
        ax.legend(fontsize=8)
        chart_index.append(save_figure(fig, folder, 27, "salary_risk_frontier", "Each path ORs failure across tested choices/macros. Dotted lines are simultaneous 95% Monte Carlo upper bounds; empty bins remain uncertified."))
    write_json(folder / "figure_index.json", chart_index)
    export_excel(folder, summary, paired, ledger, report_scope)
    return chart_index


def export_excel(folder: Path, summary: pd.DataFrame, paired: pd.DataFrame,
                 ledger: pd.DataFrame, scope: dict) -> None:
    """Application runtime export: static, typed results with inspectable source inputs."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    config = json.loads((folder / "config.json").read_text(encoding="utf-8"))
    sources = json.loads((folder / "sources.json").read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview = pd.DataFrame([{"item": "Dollar basis", "value": "2026 purchasing power; raw Parquet ledgers retain nominal accounting dollars"}, {"item": "Report type", "value": "Saved simulation snapshot; edit Python configuration and rerun to recalculate"},
                             {"item": "Reported scenarios", "value": len(summary)},
                             {"item": "Completed scenarios", "value": scope["completed_scenarios"]},
                             {"item": "Coverage", "value": scope["selection"]},
                             {"item": "Interpretation", "value": scope["notice"]}])
    tables = {"Overview": overview, "Scenarios": summary, "Amanda comparison": paired,
              "Monthly ledger 2026 USD": ledger, "Inputs": pd.DataFrame(flatten(config)), "Sources": pd.DataFrame(sources)}
    for title, filename in (("Depletion combinations", "depletion_combinations.csv"), ("Salary risk frontier", "salary_risk_frontier.csv")):
        try:
            tables[title] = pd.read_csv(folder / filename)
        except pd.errors.EmptyDataError:
            tables[title] = pd.DataFrame()
    for title, frame in tables.items():
        frame = frame.head(10_000)
        sheet = workbook.create_sheet(title)
        sheet.append(list(frame.columns) if len(frame.columns) else ["No matched cases in selection"])
        for row in frame.itertuples(index=False, name=None):
            values = []
            for value in row:
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value)
                elif isinstance(value, np.generic):
                    value = value.item()
                elif pd.isna(value):
                    value = None
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                values.append(value)
            sheet.append(values)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="233D38")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 42
        for index, column in enumerate(frame.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(55, max(18, len(column) + 3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="233D38")
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
                elif isinstance(cell.value, pd.Timestamp):
                    cell.number_format = "yyyy-mm-dd"
    temporary = folder / "report.tmp.xlsx"
    workbook.save(temporary)
    temporary.replace(folder / "report.xlsx")
